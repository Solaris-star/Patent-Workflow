from __future__ import annotations

import contextlib
import base64
import hashlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from unittest import mock

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / "patent" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import health_check_delivery_package as delivery
import init_run_manifest
import validate_facts_ledger
import validate_patent_candidates
import validate_review_status
import validate_sanitize
from generate_docx import UnsupportedMathError, generate_docx
from run_phase_gates import build_summary
from workflow_common import MODE_GATES, TITLE_MAX_CHARS, canonical_json_sha256, read_workflow_state, title_error
from validate_review_status import material_version_sha256


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def bind_sensitive_confirmation(data: dict) -> dict:
    data["confirmation_sha256"] = canonical_json_sha256({
        key: value for key, value in data.items()
        if key not in {"confirmed_by_user", "confirmed_at", "confirmation_sha256"}
    })
    return data


class ReliabilityTests(unittest.TestCase):
    def test_title_limit_is_24_characters(self) -> None:
        self.assertEqual(TITLE_MAX_CHARS, 24)
        self.assertIsNone(title_error("x" * 24))
        self.assertIsNotNone(title_error("x" * 25))

    def test_tilde_fenced_workflow_state_is_read(self) -> None:
        text = (
            "<!-- WORKFLOW_STATE_JSON_BEGIN -->\n~~~json\n"
            '{"stage_history":[{"stage":"draft"}],"waiting_for_user":true}\n'
            "~~~\n<!-- WORKFLOW_STATE_JSON_END -->"
        )
        state = read_workflow_state(text)
        self.assertEqual(state["stage_history"][0]["stage"], "draft")
        self.assertTrue(state["waiting_for_user"])

    def test_legacy_manifest_setter_does_not_consume_the_next_blank_field(self) -> None:
        text = "- output_dir:\n- final_title:\n- figure_delivery_mode: mermaid_only\n"
        updated = init_run_manifest._set_line(text, "output_dir", "C:/synthetic/output")
        self.assertIn("- output_dir: C:/synthetic/output\n", updated)
        self.assertIn("- final_title:\n", updated)
        self.assertIn("- figure_delivery_mode: mermaid_only\n", updated)

    def test_checks_passed_is_separate_from_workflow_completion(self) -> None:
        summary = build_summary(
            "all",
            [
                {"gate": "research", "status": "passed", "runs": [{"passed": True}]},
                {"gate": "deliver", "status": "skipped", "runs": []},
            ],
            review_completed=True,
            revision_validation="not_required",
            workspace=Path(".").resolve(),
        )
        self.assertTrue(summary["checks_passed"])
        self.assertTrue(summary["passed"])
        self.assertFalse(summary["workflow_complete"])
        self.assertEqual(summary["skipped"], ["deliver"])

    def test_build_summary_tolerates_unhashable_revision_result(self) -> None:
        summary = build_summary(
            "all",
            [{"gate": "review", "status": "failed", "revision_validation": [], "runs": []}],
            review_completed=False,
            revision_validation="not_run",
            workspace=Path(".").resolve(),
        )
        self.assertEqual(summary["not_run"], [])

    def test_authorized_no_revision_skip_can_still_complete_workflow(self) -> None:
        gates = ["research", "prior-art", "draft", "review", "deliver"]
        results = [
            {
                "gate": gate, "status": "passed", "runs": [{"passed": True}],
                **({"revision_validation": "not_required"} if gate == "review" else {}),
            }
            for gate in gates
        ]
        summary = build_summary(
            "all", results, review_completed=True, revision_validation="not_required",
            workspace=Path(".").resolve(),
        )
        self.assertTrue(summary["checks_passed"])
        self.assertTrue(summary["workflow_complete"])
        self.assertEqual(summary["skipped"], ["review:revision_validation"])

    def test_pending_revision_is_not_treated_as_a_no_change_skip(self) -> None:
        gates = ["research", "prior-art", "draft", "review", "deliver"]
        results = [
            {
                "gate": gate,
                "status": "passed",
                "workflow_complete": gate != "review",
                "runs": [{"passed": True, "workflow_complete": gate != "review"}],
                **({"revision_validation": "pending"} if gate == "review" else {}),
            }
            for gate in gates
        ]
        summary = build_summary(
            "all", results, review_completed=True, revision_validation="pending",
            workspace=Path(".").resolve(),
        )
        self.assertTrue(summary["checks_passed"])
        self.assertTrue(summary["review_completed"])
        self.assertEqual(summary["revision_validation"], "pending")
        self.assertFalse(summary["workflow_complete"])
        self.assertIn("review:revision_validation", summary["not_run"])

    def test_all_without_delivery_arguments_skips_delivery_without_claiming_completion(self) -> None:
        import run_phase_gates

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest = root / "run_manifest.md"
            manifest.write_text("- run_id: SYNTHETIC\n", encoding="utf-8")

            def fake_run_gate(gate, *_args):
                if gate == "deliver":
                    return [], "deliver requires output arguments", {}
                output = json.dumps({"review_completed": True}) if gate == "review" else "{}"
                return [{
                    "cmd": ["synthetic"], "exitCode": 0, "stdout": output,
                    "stderr": "", "passed": True, "checks_passed": True,
                    "workflow_complete": True, "not_run": [], "skipped": [],
                }], None, {"revision_validation": "not_required"} if gate == "review" else {}

            stdout = io.StringIO()
            with mock.patch.object(run_phase_gates, "_run_gate", side_effect=fake_run_gate), \
                    mock.patch.object(sys, "argv", [
                        "run_phase_gates.py", "--gate", "all", "--workspace", str(root),
                        "--manifest", str(manifest),
                    ]), contextlib.redirect_stdout(stdout):
                code = run_phase_gates.main()
            summary = json.loads(stdout.getvalue())
            self.assertEqual(code, 0)
            self.assertTrue(summary["checks_passed"])
            self.assertTrue(summary["passed"])
            self.assertFalse(summary["workflow_complete"])
            self.assertEqual(
                summary["skipped"], ["deliver", "review:revision_validation"]
            )

    def test_all_gate_sequences_follow_each_workflow_mode(self) -> None:
        import run_phase_gates

        for mode, stages in MODE_GATES.items():
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                manifest = root / "run_manifest.md"
                manifest.write_text(f"- workflow_mode: {mode}\n", encoding="utf-8")

                def fake_run_gate(gate, *_args):
                    result = {"workflow_complete": True}
                    details = {}
                    if gate == "review":
                        result.update({
                            "review_completed": True,
                            "revision_validation": "not_required",
                        })
                        details["revision_validation"] = "not_required"
                    return [{
                        "cmd": ["synthetic"], "exitCode": 0, "stdout": json.dumps(result),
                        "stderr": "", "passed": True, "checks_passed": True,
                        "workflow_complete": True, "not_run": [], "skipped": [],
                    }], None, details

                stdout = io.StringIO()
                with mock.patch.object(run_phase_gates, "_run_gate", side_effect=fake_run_gate), \
                        mock.patch.object(sys, "argv", [
                            "run_phase_gates.py", "--gate", "all", "--workspace", str(root),
                            "--manifest", str(manifest),
                        ]), contextlib.redirect_stdout(stdout):
                    code = run_phase_gates.main()
                summary = json.loads(stdout.getvalue())
                self.assertEqual(code, 0)
                self.assertEqual(summary["required_gates"], [gate for gate, _ in stages])
                self.assertEqual([row["gate"] for row in summary["gateResults"]], [gate for gate, _ in stages])
                self.assertTrue(summary["workflow_complete"])

    def _review_fixture(self, root: Path) -> tuple[dict, Path, Path]:
        md = root / "delivery" / "final.md"
        md.parent.mkdir(parents=True)
        md.write_text("# synthetic title\nSynthetic body.\n", encoding="utf-8")
        facts = root / "artifacts" / "draft" / "facts_ledger.json"
        write_json(facts, {"ledger_type": "facts_ledger"})
        consistency = root / "artifacts" / "audit" / "consistency.json"
        ipr = root / "artifacts" / "audit" / "ipr.json"
        hashes = {str(path.resolve()): hashlib.sha256(path.read_bytes()).hexdigest() for path in (md, facts)}
        normalized_hashes = {str(Path(path).resolve()).casefold(): digest for path, digest in hashes.items()}
        version = material_version_sha256(normalized_hashes)
        reviewed = {
            "final_markdown_path": str(md.resolve()),
            "files": hashes,
            "version_sha256": version,
        }

        def make_report(path: Path, kind: str) -> dict:
            return {
                "doc_type": f"{kind}_report",
                "report_type": kind,
                "review_id": f"REVIEW-SYN-{kind.upper()}",
                "completed_at": "2026-10-08T00:00:00Z",
                "scope": "Synthetic review fixture only.",
                "reviewed_materials": reviewed,
                "issues": [],
            }

        consistency_report = make_report(consistency, "consistency_review")
        ipr_report = make_report(ipr, "ipr_review")
        write_json(consistency, consistency_report)
        write_json(ipr, ipr_report)
        data = {
            "doc_type": "review_status",
            "review_completed": True,
            "consistency_review": {
                "status": "completed", "report_type": "consistency_review",
                "report_path": str(consistency), "report_sha256": hashlib.sha256(consistency.read_bytes()).hexdigest(),
            },
            "ipr_review": {
                "status": "completed", "report_type": "ipr_review",
                "report_path": str(ipr), "report_sha256": hashlib.sha256(ipr.read_bytes()).hexdigest(),
            },
            "reviewed_materials": reviewed,
            "issues": [],
            "revision_validation": "not_required",
        }
        return data, md, facts

    def _rewrite_review_reports(self, root: Path, data: dict, md: Path, facts: Path) -> None:
        hashes = {str(path.resolve()): hashlib.sha256(path.read_bytes()).hexdigest() for path in (md, facts)}
        normalized_hashes = {str(Path(path).resolve()).casefold(): digest for path, digest in hashes.items()}
        reviewed = data["reviewed_materials"]
        reviewed["files"] = hashes
        reviewed["version_sha256"] = material_version_sha256(normalized_hashes)
        for key in ("consistency_review", "ipr_review"):
            review = data[key]
            report_path = Path(review["report_path"])
            kind = review["report_type"]
            report = {
                "doc_type": f"{kind}_report", "report_type": kind,
                "review_id": f"REVIEW-SYN-{kind.upper()}",
                "completed_at": "2026-10-08T00:00:00Z",
                "scope": "Synthetic review fixture only.",
                "reviewed_materials": reviewed,
                "issues": data["issues"],
            }
            write_json(report_path, report)
            review["report_sha256"] = hashlib.sha256(report_path.read_bytes()).hexdigest()

    def test_review_report_presence_alone_does_not_complete_review(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data, _, _ = self._review_fixture(root)
            data["review_completed"] = False
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertTrue(errors)
            self.assertFalse(state["review_completed"])

    def test_pending_revision_keeps_review_completion_separate_from_workflow_completion(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data, _, _ = self._review_fixture(root)
            data["revision_validation"] = "pending"
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertEqual(errors, [])
            self.assertTrue(state["review_completed"])
            self.assertFalse(state["workflow_complete"])

    def test_review_reports_must_be_structured_current_and_bound_to_materials(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data, _md, _facts = self._review_fixture(root / "placeholder")
            for key in ("consistency_review", "ipr_review"):
                report_path = Path(data[key]["report_path"])
                report_path.write_text("synthetic nonempty placeholder", encoding="utf-8")
                data[key]["report_sha256"] = hashlib.sha256(report_path.read_bytes()).hexdigest()
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertTrue(any("structured JSON" in error for error in errors))
            self.assertFalse(state["review_completed"])

            data, _md, _facts = self._review_fixture(root / "fresh")
            report_path = Path(data["consistency_review"]["report_path"])
            report_path.write_text(report_path.read_text(encoding="utf-8") + " ", encoding="utf-8")
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertTrue(any("report hash" in error for error in errors))
            self.assertFalse(state["review_fresh"])

    def test_review_detects_changed_hash_and_requires_concrete_high_issue_waiver(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data, md, _ = self._review_fixture(root)
            facts = root / "artifacts" / "draft" / "facts_ledger.json"
            data["issues"] = [{
                "issue_id": "ISSUE-SYN-1", "severity": "high", "disposition": "open",
                "summary": "Synthetic unresolved risk.",
            }]
            self._rewrite_review_reports(root, data, md, facts)
            errors, _ = validate_review_status.validate_review_status(data, workspace=root)
            self.assertTrue(any("waiver" in error or "high-severity" in error for error in errors))

            data["issues"][0]["disposition"] = "waived"
            data["user_waivers"] = [{
                "issue_ids": ["ISSUE-SYN-1"], "decision": "accept_specific_risk",
                "confirmed_by_user": True, "confirmed_at": "2026-10-08T00:00:00Z",
                "approval_reference": "APPROVAL-SYN-1",
                "material_version_sha256": data["reviewed_materials"]["version_sha256"],
                "scope": {
                    "kind": "issue_set", "issue_ids": ["ISSUE-SYN-1"],
                    "limitations": ["Applies only to the named unresolved issue in this exact synthetic review version."],
                },
                "accepted_risks": {"ISSUE-SYN-1": "The review will retain this identified synthetic issue as unresolved."},
            }]
            self._rewrite_review_reports(root, data, md, facts)
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertFalse(errors)
            self.assertTrue(state["review_completed"])
            self.assertFalse(state["semantic_review_assessed"])
            self.assertFalse(state["identity_authenticated"])

            md.write_text("# changed after review\n", encoding="utf-8")
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertTrue(any("changed" in error for error in errors))
            self.assertFalse(state["review_fresh"])

    def test_post_fix_report_must_match_latest_hashes_and_issue_dispositions(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data, md, facts = self._review_fixture(root)
            data["issues"] = [{
                "issue_id": "ISSUE-SYN-2", "severity": "medium", "disposition": "resolved",
                "summary": "Synthetic resolved issue.",
            }]
            data["revision_validation"] = "passed"
            self._rewrite_review_reports(root, data, md, facts)
            report_path = root / "artifacts" / "revision" / "post_fix.json"
            hashes = {
                str(md): hashlib.sha256(md.read_bytes()).hexdigest(),
                str(facts): hashlib.sha256(facts.read_bytes()).hexdigest(),
            }
            write_json(report_path, {
                "checks_passed": True,
                "material_hashes": hashes,
                "issue_outcomes": {"ISSUE-SYN-2": "resolved"},
            })
            data["post_fix_report_path"] = str(report_path)
            errors, state = validate_review_status.validate_review_status(
                data, workspace=root, require_revision=True
            )
            self.assertFalse(errors)
            self.assertTrue(state["review_completed"])

            tampered = json.loads(report_path.read_text(encoding="utf-8"))
            tampered["material_hashes"][str(md)] = "0" * 64
            write_json(report_path, tampered)
            errors, _ = validate_review_status.validate_review_status(
                data, workspace=root, require_revision=True
            )
            self.assertTrue(any("hashes" in error for error in errors))

    def test_generic_high_issue_waivers_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data, md, facts = self._review_fixture(root)
            data["issues"] = [{
                "issue_id": "ISSUE-SYN-GENERIC", "severity": "high", "disposition": "waived",
                "summary": "Synthetic unresolved issue.",
            }]
            data["user_waivers"] = [{
                "issue_ids": ["ISSUE-SYN-GENERIC"], "decision": "yes", "confirmed_by_user": True,
                "confirmed_at": "2026-10-08T00:00:00Z", "approval_reference": "APPROVAL-SYN-GENERIC",
                "material_version_sha256": data["reviewed_materials"]["version_sha256"],
                "scope": {"kind": "all", "issue_ids": ["ISSUE-SYN-GENERIC"], "limitations": ["all"]},
                "accepted_risks": {"ISSUE-SYN-GENERIC": "yes"},
            }]
            self._rewrite_review_reports(root, data, md, facts)
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertTrue(any("decision" in error for error in errors))
            self.assertTrue(any("limitations" in error for error in errors))
            self.assertTrue(any("specific accepted-risk reason" in error for error in errors))
            self.assertFalse(state["review_completed"])

    def test_high_severity_waivers_are_required_for_each_issue_id(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data, md, facts = self._review_fixture(root)
            issue_ids = ["ISSUE-SYN-HIGH-1", "ISSUE-SYN-HIGH-2"]
            data["issues"] = [{
                "issue_id": issue_id, "severity": "high", "disposition": "waived",
                "summary": "Synthetic issue-specific risk.",
            } for issue_id in issue_ids]
            self._rewrite_review_reports(root, data, md, facts)

            def waiver(issue_id: str) -> dict:
                return {
                    "issue_ids": [issue_id], "decision": "accept_specific_risk",
                    "confirmed_by_user": True, "confirmed_at": "2026-10-08T00:00:00Z",
                    "approval_reference": f"APPROVAL-{issue_id}",
                    "material_version_sha256": data["reviewed_materials"]["version_sha256"],
                    "scope": {
                        "kind": "issue_set", "issue_ids": [issue_id],
                        "limitations": [f"Applies only to {issue_id} in this synthetic review version."],
                    },
                    "accepted_risks": {issue_id: f"Retain only the named synthetic risk {issue_id}."},
                }

            data["user_waivers"] = [waiver(issue_ids[0])]
            errors, _ = validate_review_status.validate_review_status(data, workspace=root)
            self.assertTrue(any(issue_ids[1] in error and "waiver" in error for error in errors))

            data["user_waivers"].append(waiver(issue_ids[1]))
            errors, state = validate_review_status.validate_review_status(data, workspace=root)
            self.assertEqual(errors, [])
            self.assertTrue(state["review_completed"])

    def test_unconfirmed_sensitive_map_is_not_scanned_or_echoed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            marker = "synthetic-marker-secret"
            sensitive_map = root / "map.json"
            data = bind_sensitive_confirmation({
                "map_type": "sensitive_map",
                "confirmed_by_user": False,
                "confirmation_scope": "synthetic test only",
                "confirmed_at": "2026-10-08T00:00:00Z",
                "entries": [{"id": "S-1", "term": marker, "aliases": []}],
            })
            write_json(sensitive_map, data)
            target = root / "target.txt"
            target.write_text(marker, encoding="utf-8")
            output = io.StringIO()
            with mock.patch.object(sys, "argv", [
                "validate_sanitize.py", "--map", str(sensitive_map), "--scan-dir", str(root),
            ]), contextlib.redirect_stdout(output):
                result = validate_sanitize.main()
            rendered = output.getvalue()
            self.assertEqual(result, 2)
            self.assertNotIn(marker, rendered)
            self.assertIn('"mapHits": 0', rendered)

    def test_validate_only_checks_required_sensitive_map_entries(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            sensitive_map = root / "map.json"
            data = bind_sensitive_confirmation({
                "map_type": "sensitive_map",
                "confirmed_by_user": True,
                "confirmation_scope": "synthetic test only",
                "confirmed_at": "2026-10-08T00:00:00Z",
                "entries": [{"id": "S-1", "term": "", "aliases": []}],
            })
            write_json(sensitive_map, data)
            output = io.StringIO()
            with mock.patch.object(sys, "argv", ["validate_sanitize.py", "--map", str(sensitive_map), "--validate-only"]), contextlib.redirect_stdout(output):
                result = validate_sanitize.main()
            self.assertEqual(result, 2)
            self.assertIn("missing its literal term", output.getvalue())

    def test_sensitive_scan_does_not_echo_private_paths(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            map_path = root / "map.json"
            data = bind_sensitive_confirmation({
                "map_type": "sensitive_map",
                "confirmed_by_user": True,
                "confirmation_scope": "synthetic test only",
                "confirmed_at": "2026-10-08T00:00:00Z",
                "entries": [{"id": "S-1", "term": "synthetic marker", "aliases": []}],
            })
            write_json(map_path, data)
            private_path = root / "private-local-marker"
            output = io.StringIO()
            with mock.patch.object(sys, "argv", [
                "validate_sanitize.py", "--map", str(map_path), "--scan-dir", str(private_path),
            ]), contextlib.redirect_stdout(output):
                result = validate_sanitize.main()
            self.assertEqual(result, 2)
            self.assertNotIn(str(private_path), output.getvalue())

    def test_sensitive_map_confirmation_binds_entries_and_scope(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            marker = "synthetic-private-term"
            map_path = root / "map.json"
            data = bind_sensitive_confirmation({
                "map_type": "sensitive_map", "confirmed_by_user": True,
                "confirmation_scope": "Only this synthetic example delivery.",
                "confirmed_at": "2026-10-08T00:00:00Z",
                "entries": [{"id": "S-1", "term": marker, "aliases": []}],
            })
            write_json(map_path, data)
            target = root / "target.txt"
            target.write_text(marker, encoding="utf-8")
            data["confirmation_scope"] = "A broader changed scope."
            write_json(map_path, data)
            output = io.StringIO()
            with mock.patch.object(sys, "argv", [
                "validate_sanitize.py", "--map", str(map_path), "--scan-dir", str(root),
            ]), contextlib.redirect_stdout(output):
                result = validate_sanitize.main()
            self.assertEqual(result, 2)
            report = json.loads(output.getvalue())
            self.assertFalse(report["confirmation_binding_valid"])
            self.assertEqual(report["counts"]["mapHits"], 0)
            self.assertNotIn(marker, output.getvalue())

    def test_candidate_validation_has_no_default_count_or_recent_date_quota(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidates = root / "candidates.json"
            write_json(candidates, [{
                "title": "Synthetic old candidate",
                "publicationNumber": "CN-SYN-1",
                "filingDate": "2001-01-01",
                "abstract": "synthetic record",
            }])
            output = io.StringIO()
            with mock.patch.object(sys, "argv", [
                "validate_patent_candidates.py", str(candidates), "--relevance-threshold", "0",
            ]), contextlib.redirect_stdout(output):
                with self.assertRaises(SystemExit) as raised:
                    validate_patent_candidates.main()
            self.assertEqual(raised.exception.code, 0)
            summary = json.loads(output.getvalue())
            self.assertEqual(summary["thresholds"]["minCount"], 0)
            self.assertTrue(summary["gates"]["freshnessPassed"])

            output = io.StringIO()
            with mock.patch.object(sys, "argv", [
                "validate_patent_candidates.py", str(candidates), "--relevance-threshold", "0",
                "--fresh-years", "1", "--min-count", "1",
            ]), contextlib.redirect_stdout(output):
                with self.assertRaises(SystemExit) as raised:
                    validate_patent_candidates.main()
            self.assertEqual(raised.exception.code, 2)

    def test_mermaid_only_ledger_accepts_no_image_and_requires_all_five_parts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            parts = root / "artifacts" / "draft"
            parts.mkdir(parents=True)
            for index in range(1, 6):
                extra = " FIG-1 Synthetic figure." if index == 4 else ""
                (parts / f"part_{index:02d}_synthetic.md").write_text(
                    f"# Part {index}\nSynthetic section content {index}.{extra}\n", encoding="utf-8"
                )
            mmd = parts / "figure_01.mmd"
            mmd.write_text("flowchart TD\n  A-->B\n", encoding="utf-8")
            ledger = {
                "ledger_type": "facts_ledger",
                "terminology": [{"term": "synthetic term", "definition": "synthetic definition"}],
                "constraints_and_effects": [{"constraint": "synthetic", "effect": "synthetic"}],
                "source_registry": [{
                    "source_id": "SRC-FIG-1", "source_type": "figure", "figure_id": "FIG-1",
                    "path": str(mmd.relative_to(root)), "sha256": hashlib.sha256(mmd.read_bytes()).hexdigest(),
                }],
                "feature_registry": [{
                    "feature_id": "F-SYN-1", "statement": "synthetic feature",
                    "evidence_kind": "implemented_fact", "status": "confirmed",
                    "source_ids": ["SRC-FIG-1"], "evidence_ids": [], "paragraph_ids": [], "figure_ids": ["FIG-1"],
                }],
                "figure_registry": [{
                    "figure_id": "FIG-1", "caption": "Synthetic figure",
                    "artifacts": {"mmd": str(mmd.relative_to(root))},
                    "mermaid_source_embedded_in_docx": True,
                }],
            }
            write_json(root / "artifacts" / "draft" / "facts_ledger.json", ledger)
            manifest = root / "artifacts" / "run_manifest.md"
            manifest.parent.mkdir(parents=True, exist_ok=True)
            manifest.write_text(
                "- `final_title`: " + "x" * 24 + "\n- `figure_delivery_mode`: mermaid_only\n",
                encoding="utf-8",
            )
            errors = validate_facts_ledger.validate_ledger(
                ledger, base_dir=root, figure_delivery_mode="mermaid_only", require_main_parts=True,
                require_docx_visible_mermaid=True,
            )
            self.assertEqual(errors, [])
            other_mmd = parts / "other_figure.mmd"
            other_mmd.write_text("flowchart TD\n  C-->D\n", encoding="utf-8")
            ledger["figure_registry"][0]["artifacts"]["mmd"] = str(other_mmd.relative_to(root))
            errors = validate_facts_ledger.validate_ledger(
                ledger, base_dir=root, figure_delivery_mode="mermaid_only", require_main_parts=True,
                require_docx_visible_mermaid=True,
            )
            self.assertTrue(any("must match its figure source path" in error for error in errors))
            ledger["figure_registry"][0]["artifacts"]["mmd"] = str(mmd.relative_to(root))
            (parts / "part_03_synthetic.md").unlink()
            errors = validate_facts_ledger.validate_ledger(
                ledger, base_dir=root, figure_delivery_mode="mermaid_only", require_main_parts=True,
                require_docx_visible_mermaid=True,
            )
            self.assertTrue(any("part_03" in error for error in errors))

    @unittest.skipIf(__import__("importlib").util.find_spec("docx") is None, "python-docx not installed")
    def test_docx_export_is_repeatable_and_preserves_source_markdown(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "final.md"
            output = root / "final.docx"
            source.write_text("# Synthetic title\nSynthetic body.\n", encoding="utf-8")
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            first = generate_docx(source, output)
            with self.assertRaises(FileExistsError):
                generate_docx(source, output)
            second = generate_docx(source, output, overwrite=True)
            self.assertTrue(source.is_file())
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
            self.assertEqual(first["output_path"], str(output.resolve()))
            self.assertEqual(second["output_path"], str(output.resolve()))
            self.assertTrue(output.is_file())

    @unittest.skipIf(__import__("importlib").util.find_spec("docx") is None, "python-docx not installed")
    def test_docx_preserves_chinese_bold_tables_and_common_math_as_omml(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            from docx import Document

            root = Path(tmp)
            source = root / "synthetic-rich.md"
            output = root / "synthetic-rich.docx"
            markdown = (
                "# 软件专利协作流程\n"
                "正文**加粗中文**。公式 $x^2$ 与 $\\frac{\\alpha+\\beta}{\\gamma}$ 以及 $A|B$。\n\n"
                "| 类型 | 内容 |\n| --- | --- |\n| 公式 | $\\alpha + \\beta$ |\n| 特征 | **关键步骤** |\n| 分隔 | left\\|right |\n"
            )
            source.write_text(markdown, encoding="utf-8")
            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            generated = generate_docx(source, output)
            self.assertEqual(generated["status"], "generated")
            self.assertTrue(source.is_file())
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_hash)

            document = Document(output)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 0).text, "公式")
            self.assertEqual(document.tables[0].cell(1, 1).text, "")
            self.assertEqual(document.tables[0].cell(2, 1).text, "关键步骤")
            self.assertTrue(document.tables[0].cell(2, 1).paragraphs[0].runs[0].bold)
            self.assertEqual(document.tables[0].cell(3, 1).text, "left|right")

            namespaces = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            }
            with zipfile.ZipFile(output) as package:
                xml_root = ET.fromstring(package.read("word/document.xml"))
            math_nodes = xml_root.findall(".//m:oMath", namespaces)
            self.assertEqual(len(math_nodes), 4)
            self.assertTrue(xml_root.findall(".//m:sSup", namespaces))
            self.assertTrue(xml_root.findall(".//m:f", namespaces))
            math_tokens = [
                "".join(node.text or "" for node in formula.findall(".//m:t", namespaces))
                for formula in math_nodes
            ]
            self.assertEqual(math_tokens, ["x2", "α+βγ", "A|B", "α+β"])
            table_rows = xml_root.findall(".//w:tbl/w:tr", namespaces)
            table_formula = table_rows[1].findall(".//m:oMath", namespaces)
            self.assertEqual(len(table_formula), 1)
            self.assertEqual(
                "".join(node.text or "" for node in table_formula[0].findall(".//m:t", namespaces)),
                "α+β",
            )
            bold_texts = []
            for run in xml_root.findall(".//w:r", namespaces):
                if run.find("w:rPr/w:b", namespaces) is not None:
                    bold_texts.append(
                        "".join(node.text or "" for node in run.findall(".//w:t", namespaces))
                    )
            self.assertTrue(any("加粗中文" in text for text in bold_texts))
            visible_units = delivery._markdown_visible_text_units(
                "Equation $A|B$\n|Header|Cell|\n|---|---|\n|value|result|"
            )
            self.assertIn("Equation", visible_units)
            self.assertNotIn("$A", visible_units)
            self.assertIn("Header", visible_units)
            self.assertTrue(delivery._markdown_text_present("简单公式：$x^2$。", "简单公式：。"))
            self.assertFalse(delivery._markdown_text_present("A $x^2$ B", "AB"))
            self.assertTrue(delivery._markdown_text_present("A$x^2$B", "AB"))

    @unittest.skipIf(__import__("importlib").util.find_spec("docx") is None, "python-docx not installed")
    def test_unsupported_inline_math_is_rejected_without_replacing_old_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "unsupported.md"
            output = root / "existing.docx"
            unsupported_cases = (
                ("unknown command", "Unsupported $\\unknown{x}$."),
                ("square root", "Unsupported $\\sqrt{x}$."),
                ("matrix", "Unsupported $\\begin{matrix}x\\end{matrix}$."),
                ("unpaired delimiter", "Unsupported $x."),
                ("display math", "Unsupported $$x^2$$."),
            )
            old_output = b"previous synthetic output"
            for label, body in unsupported_cases:
                with self.subTest(math_case=label):
                    source.write_text(f"# Synthetic title\n{body}\n", encoding="utf-8")
                    output.write_bytes(old_output)
                    original_source = source.read_bytes()
                    with self.assertRaises(UnsupportedMathError):
                        generate_docx(source, output, overwrite=True)
                    self.assertEqual(output.read_bytes(), old_output)
                    self.assertEqual(source.read_bytes(), original_source)

    @unittest.skipIf(__import__("importlib").util.find_spec("docx") is None, "python-docx not installed")
    def test_mermaid_delivery_passes_structural_checks_and_render_not_run_is_explicit(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            title = "Synthetic patent"
            review_data, md_path, facts = self._review_fixture(root)
            deliver_dir = md_path.parent
            mmd_path = deliver_dir / "figure_01.mmd"
            fence = chr(96) * 3
            md_path.write_text(
                f"# {title}\nSynthetic body [[E-SYN-1]] with **bold text** and $x^2$ plus $\\frac{{\\alpha+\\beta}}{{\\gamma}}$. 简单公式：$x^2$。\n\n"
                "| Feature | Evidence | Formula | Bold | Pipe |\n| --- | --- | --- | --- | --- |\n"
                "| F-SYN-1 | [[E-SYN-1]] | $\\alpha + \\beta$ | **加粗单元格** | left\\|right |\n\n"
                f"{fence}mermaid\nflowchart TD\n  A-->B\n{fence}\n",
                encoding="utf-8",
            )
            mmd_path.write_text("flowchart TD\n  A-->B\n", encoding="utf-8")
            docx_path = deliver_dir / f"{title}{delivery.DOCX_SUFFIX}"
            self.assertIsNotNone(generate_docx(md_path, docx_path))
            write_json(facts, {
                "ledger_type": "facts_ledger",
                "feature_registry": [{"feature_id": "F-SYN-1", "evidence_ids": ["E-SYN-1"]}],
                "figure_registry": [{
                    "figure_id": "FIG-SYN-1", "caption": "Synthetic figure",
                    "artifacts": {"mmd": "figure_01.mmd"},
                }],
            })
            consistency = Path(review_data["consistency_review"]["report_path"])
            ipr = Path(review_data["ipr_review"]["report_path"])
            manifest = root / "artifacts" / "run_manifest.md"
            manifest.write_text(
                f"- output_dir: {deliver_dir.resolve()}\n- final_title: {title}\n"
                "- figure_delivery_mode: mermaid_only\n",
                encoding="utf-8",
            )
            review = root / "artifacts" / "audit" / "review_status.json"
            self._rewrite_review_reports(root, review_data, md_path, facts)
            write_json(review, review_data)
            with mock.patch.object(delivery, "_render_docx", return_value={"status": "not_run", "reason": "synthetic renderer unavailable"}):
                result = delivery.check_delivery(
                    deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                    facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                    review_status_path=review, manifest_path=manifest, base_dir=root,
                    figure_delivery_mode="mermaid_only",
                )
            self.assertTrue(result["checks_passed"], result["errors"])
            self.assertFalse(result["workflow_complete"])
            self.assertIn("docx_render", result["not_run"])
            self.assertEqual(result["render"]["status"], "not_run")
            math_check = next(check for check in result["checks"] if check["name"] == "DOCX inline math matches Markdown equations")
            self.assertEqual(math_check["result"], "pass", math_check["details"])
            text_check = next(check for check in result["checks"] if check["name"] == "DOCX retains visible Markdown text and table cells")
            self.assertEqual(text_check["result"], "pass", text_check["details"])
            generated_docx = Document(docx_path)
            self.assertEqual(len(generated_docx.tables), 1)
            self.assertEqual(generated_docx.tables[0].cell(1, 0).text, "F-SYN-1")

            extracted_text, embedded_hashes, embedded_math = delivery._docx_text_and_image_refs(docx_path)
            self.assertIn("简单公式：。", extracted_text)
            missing_one_character = extracted_text.replace("Synthetic body", "Synthetic bod", 1)
            with mock.patch.object(
                delivery,
                "_docx_text_and_image_refs",
                return_value=(missing_one_character, embedded_hashes, embedded_math),
            ), mock.patch.object(
                delivery, "_render_docx", return_value={"status": "not_run", "reason": "synthetic renderer unavailable"},
            ):
                missing_character = delivery.check_delivery(
                    deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                    facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                    review_status_path=review, manifest_path=manifest, base_dir=root,
                    figure_delivery_mode="mermaid_only",
                )
            self.assertFalse(missing_character["checks_passed"])
            text_check = next(check for check in missing_character["checks"] if check["name"] == "DOCX retains visible Markdown text and table cells")
            self.assertEqual(text_check["result"], "fail")

            missing_table_cell = extracted_text.replace("F-SYN-1", "F-SYN", 1)
            with mock.patch.object(
                delivery,
                "_docx_text_and_image_refs",
                return_value=(missing_table_cell, embedded_hashes, embedded_math),
            ), mock.patch.object(
                delivery, "_render_docx", return_value={"status": "not_run", "reason": "synthetic renderer unavailable"},
            ):
                omitted_table_cell = delivery.check_delivery(
                    deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                    facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                    review_status_path=review, manifest_path=manifest, base_dir=root,
                    figure_delivery_mode="mermaid_only",
                )
            self.assertFalse(omitted_table_cell["checks_passed"])
            text_check = next(check for check in omitted_table_cell["checks"] if check["name"] == "DOCX retains visible Markdown text and table cells")
            self.assertEqual(text_check["result"], "fail")

            render_failure_cases = (
                ("converter exit failure", mock.Mock(return_value=type("Result", (), {"returncode": 1})())),
                ("successful exit without PDF", mock.Mock(return_value=type("Result", (), {"returncode": 0})())),
                ("render timeout", mock.Mock(side_effect=delivery.subprocess.TimeoutExpired("soffice", 90))),
            )
            for label, conversion in render_failure_cases:
                with self.subTest(render_case=label):
                    with mock.patch.object(delivery.shutil, "which", return_value="soffice"), \
                         mock.patch.object(delivery.subprocess, "run", conversion):
                        render_failed = delivery.check_delivery(
                            deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                            facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                            review_status_path=review, manifest_path=manifest, base_dir=root,
                            figure_delivery_mode="mermaid_only",
                        )
                    self.assertEqual(render_failed["render"]["status"], "failed")
                    self.assertFalse(render_failed["checks_passed"])
                    self.assertFalse(render_failed["workflow_complete"])
                    self.assertIn("DOCX render health", render_failed["errors"])
                    render_check = next(check for check in render_failed["checks"] if check["name"] == "DOCX render health")
                    self.assertEqual(render_check["result"], "fail")

            renderer_status_cases = (("failed", "fail"), ("fail", "fail"), ("error", "fail"), ("mystery", "fail"))
            for renderer_status, expected_check_status in renderer_status_cases:
                with self.subTest(renderer_status=renderer_status):
                    with mock.patch.object(
                        delivery, "_render_docx", return_value={"status": renderer_status, "reason": "synthetic renderer status"},
                    ):
                        status_failure = delivery.check_delivery(
                            deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                            facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                            review_status_path=review, manifest_path=manifest, base_dir=root,
                            figure_delivery_mode="mermaid_only",
                        )
                    self.assertFalse(status_failure["checks_passed"])
                    self.assertFalse(status_failure["workflow_complete"])
                    self.assertIn("DOCX render health", status_failure["errors"])
                    status_check = next(check for check in status_failure["checks"] if check["name"] == "DOCX render health")
                    self.assertEqual(status_check["result"], expected_check_status)

            with mock.patch.object(
                delivery, "_render_docx", return_value={"status": "mystery", "reason": "synthetic unknown renderer state"},
            ):
                unknown_render = delivery.check_delivery(
                    deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                    facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                    review_status_path=review, manifest_path=manifest, base_dir=root,
                    figure_delivery_mode="mermaid_only",
                )
            self.assertFalse(unknown_render["checks_passed"])
            self.assertFalse(unknown_render["workflow_complete"])
            self.assertIn("DOCX render health", unknown_render["errors"])
            unknown_check = next(check for check in unknown_render["checks"] if check["name"] == "DOCX render health")
            self.assertEqual(unknown_check["result"], "fail")
            self.assertIn("unrecognized check result status", unknown_check["details"])

            manifest.write_text(
                f"- output_dir: {root / 'wrong-delivery'}\n- final_title: {title}\n"
                "- figure_delivery_mode: mermaid_only\n",
                encoding="utf-8",
            )
            with mock.patch.object(delivery, "_render_docx", return_value={"status": "not_run", "reason": "synthetic renderer unavailable"}):
                mismatch = delivery.check_delivery(
                    deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                    facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                    review_status_path=review, manifest_path=manifest, base_dir=root,
                    figure_delivery_mode="mermaid_only",
                )
            self.assertFalse(mismatch["checks_passed"])
            self.assertIn("manifest output_dir matches delivery directory", mismatch["errors"])

            image_path = deliver_dir / "figure.png"
            image_path.write_bytes(base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+i3ioAAAAASUVORK5CYII="
            ))
            facts_data = json.loads(facts.read_text(encoding="utf-8"))
            facts_data["figure_registry"][0]["artifacts"]["image"] = "figure.png"
            write_json(facts, facts_data)
            manifest.write_text(
                f"- output_dir: {deliver_dir.resolve()}\n- final_title: {title}\n"
                "- figure_delivery_mode: mermaid_and_images\n",
                encoding="utf-8",
            )
            md_path.write_text(
                f"# {title}\nSynthetic body [[E-SYN-1]].\n\n{fence}mermaid\nflowchart TD\n  A-->B\n{fence}\n\n"
                "![Synthetic figure](figure.png)\n",
                encoding="utf-8",
            )
            self.assertIsNotNone(generate_docx(md_path, docx_path, overwrite=True))
            self._rewrite_review_reports(root, review_data, md_path, facts)
            write_json(review, review_data)
            with mock.patch.object(delivery, "_render_docx", return_value={"status": "not_run", "reason": "synthetic renderer unavailable"}):
                image_result = delivery.check_delivery(
                    deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                    facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                    review_status_path=review, manifest_path=manifest, base_dir=root,
                    figure_delivery_mode="mermaid_and_images",
                )
            self.assertTrue(image_result["checks_passed"], image_result["errors"])
            self.assertFalse(image_result["workflow_complete"])

            md_path.write_text(md_path.read_text(encoding="utf-8").replace("figure.png", "missing.png"), encoding="utf-8")
            self.assertIsNotNone(generate_docx(md_path, docx_path, overwrite=True))
            self._rewrite_review_reports(root, review_data, md_path, facts)
            write_json(review, review_data)
            with mock.patch.object(delivery, "_render_docx", return_value={"status": "not_run", "reason": "synthetic renderer unavailable"}):
                missing_image = delivery.check_delivery(
                    deliver_dir=deliver_dir, patent_title=title, final_markdown=md_path,
                    facts_ledger_path=facts, consistency_report=consistency, ipr_report=ipr,
                    review_status_path=review, manifest_path=manifest, base_dir=root,
                    figure_delivery_mode="mermaid_and_images",
                )
            self.assertFalse(missing_image["checks_passed"])
            self.assertIn("referenced image files exist and match configured mode", missing_image["errors"])


if __name__ == "__main__":
    unittest.main()
