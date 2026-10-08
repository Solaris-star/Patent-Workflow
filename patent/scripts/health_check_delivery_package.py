#!/usr/bin/env python3
"""Check the actual delivery directory, reviewed version, DOCX, citations, and figures."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from generate_docx import UnsupportedMathError, _split_table_row, inline_math_expressions, math_expression_text
from validate_review_status import validate_review_status
from workflow_common import DOCX_SUFFIX, TITLE_MAX_CHARS, is_within, manifest_field_text, sha256_file

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, OSError, ValueError):
        pass

IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
IMAGE_ALT_RE = re.compile(r"!\[([^\]]*)\]\([^)]+\)")
INLINE_MATH_RE = re.compile(r"(?<!\\)\$([^$]+)(?<!\\)\$")
CHECK_STATUS_ALIASES = {
    "pass": "pass",
    "passed": "pass",
    "not_run": "not_run",
    "fail": "fail",
    "failed": "fail",
    "failure": "fail",
    "error": "fail",
    "errored": "fail",
}
DOCX_NS = {
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}


def _normalized_check_status(status: object) -> str:
    if not isinstance(status, str):
        return "fail"
    key = status.strip().casefold().replace("-", "_").replace(" ", "_")
    return CHECK_STATUS_ALIASES.get(key, "fail")


def _check(checks: list[dict], name: str, status: object, details: object = "") -> None:
    normalized = _normalized_check_status(status)
    detail_text = details if isinstance(details, str) else "invalid check details"
    status_key = (
        status.strip().casefold().replace("-", "_").replace(" ", "_")
        if isinstance(status, str) else ""
    )
    if normalized == "fail" and status_key not in CHECK_STATUS_ALIASES:
        detail_text = (detail_text + "; " if detail_text else "") + "unrecognized check result status"
    checks.append({"name": name, "result": normalized, "details": detail_text})


def _manifest_value(manifest_path: Path, key: str) -> str | None:
    if not manifest_path.is_file():
        return None
    return manifest_field_text(manifest_path.read_text(encoding="utf-8"), key)


def _delivery_path(deliver_dir: Path, raw: str) -> Path | None:
    path = Path(raw)
    if not path.is_absolute():
        path = deliver_dir / path
    if not is_within(path, deliver_dir):
        return None
    return path


def _omml_math_text(node: ET.Element) -> str:
    local_name = node.tag.rsplit("}", 1)[-1]
    math_namespace = DOCX_NS["m"]
    if local_name == "t" and node.tag.startswith(f"{{{math_namespace}}}"):
        return node.text or ""
    if local_name == "d" and node.tag.startswith(f"{{{math_namespace}}}"):
        properties = node.find("m:dPr", DOCX_NS)
        opening = properties.find("m:begChr", DOCX_NS) if properties is not None else None
        closing = properties.find("m:endChr", DOCX_NS) if properties is not None else None
        expression = node.find("m:e", DOCX_NS)
        return (
            (opening.get(f"{{{math_namespace}}}val", "") if opening is not None else "")
            + (_omml_math_text(expression) if expression is not None else "")
            + (closing.get(f"{{{math_namespace}}}val", "") if closing is not None else "")
        )
    if local_name == "f" and node.tag.startswith(f"{{{math_namespace}}}"):
        numerator = node.find("m:num", DOCX_NS)
        denominator = node.find("m:den", DOCX_NS)
        return (_omml_math_text(numerator) if numerator is not None else "") + (
            _omml_math_text(denominator) if denominator is not None else ""
        )
    if local_name in {"sSub", "sSup", "sSubSup"} and node.tag.startswith(f"{{{math_namespace}}}"):
        base = node.find("m:e", DOCX_NS)
        sub = node.find("m:sub", DOCX_NS)
        sup = node.find("m:sup", DOCX_NS)
        return (
            (_omml_math_text(base) if base is not None else "")
            + (_omml_math_text(sub) if sub is not None else "")
            + (_omml_math_text(sup) if sup is not None else "")
        )
    return "".join(_omml_math_text(child) for child in node)


def _docx_text_and_image_refs(docx_path: Path) -> tuple[str, list[str], list[str]]:
    from docx import Document

    doc = Document(docx_path)
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    embedded_image_hashes: list[str] = []
    embedded_math_texts: list[str] = []
    with zipfile.ZipFile(docx_path, "r") as package:
        document_root = ET.fromstring(package.read("word/document.xml"))
        rels_root = ET.fromstring(package.read("word/_rels/document.xml.rels"))
        for formula in document_root.findall(".//m:oMath", DOCX_NS):
            formula_text = _omml_math_text(formula)
            embedded_math_texts.append(formula_text)
            paragraphs.append(formula_text)
        relations = {
            relation.attrib.get("Id", ""): relation.attrib.get("Target", "")
            for relation in rels_root.findall("rel:Relationship", DOCX_NS)
            if relation.attrib.get("Type", "").endswith("/image")
        }
        for blip in document_root.findall(".//a:blip", DOCX_NS):
            relation_id = blip.attrib.get(f"{{{DOCX_NS['r']}}}embed", "")
            target = relations.get(relation_id, "")
            member = posixpath.normpath(posixpath.join("word", target))
            if not target or member not in package.namelist():
                raise ValueError("DOCX image relationship is unresolved")
            embedded_image_hashes.append(hashlib.sha256(package.read(member)).hexdigest())
    return "\n".join(paragraphs), embedded_image_hashes, embedded_math_texts


def _markdown_visible_text_units(markdown: str) -> list[str]:
    """Extract individual human-readable Markdown units for DOCX presence checks."""
    units: list[str] = []
    lines = markdown.splitlines()
    in_fence = False
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith(chr(96) * 3):
            in_fence = not in_fence
            index += 1
            continue
        if in_fence or not line:
            index += 1
            continue

        if index + 1 < len(lines) and "|" in line and "|" in lines[index + 1]:
            cells = _split_table_row(line)
            separators = _split_table_row(lines[index + 1])
            if len(cells) == len(separators) and cells and all(
                re.fullmatch(r":?-{3,}:?", cell) for cell in separators
            ):
                units.extend(
                    visible for cell in cells if (visible := _normalize_inline_markdown_text(cell))
                )
                index += 2
                while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                    units.extend(
                        visible
                        for cell in _split_table_row(lines[index])
                        if (visible := _normalize_inline_markdown_text(cell))
                    )
                    index += 1
                continue

        line = re.sub(r"^\s{0,3}#{1,6}\s+", "", line)
        line = re.sub(r"^\s*>\s?", "", line)
        line = re.sub(r"^\s*(?:[-+*]|\d+[.)])\s+", "", line)
        line = _normalize_inline_markdown_text(line)
        if line:
            units.append(line)
        index += 1
    return units


def _normalize_inline_markdown_text(text: str) -> str:
    text = IMAGE_ALT_RE.sub(r"\1", text)
    text = re.sub(r"(?<!!)\[([^\]]+)\]\([^)]+\)", r"\1", text)
    # OMML is absent from python-docx paragraph.text. Remove only the formula
    # itself: surrounding source whitespace remains significant.
    text = INLINE_MATH_RE.sub("", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return re.sub(r"\s+", " ", text).strip()


def _markdown_text_present(markdown: str, docx_text: str) -> bool:
    normalized_docx = re.sub(r"\s+", " ", docx_text).casefold()
    return all(
        re.sub(r"\s+", " ", unit).casefold() in normalized_docx
        for unit in _markdown_visible_text_units(markdown)
    )


def _render_docx(docx_path: Path) -> dict:
    renderer = shutil.which("soffice") or shutil.which("libreoffice")
    if not renderer:
        return {"status": "not_run", "reason": "LibreOffice renderer is not installed"}
    with tempfile.TemporaryDirectory(prefix="patent-docx-render-") as temp_dir:
        try:
            result = subprocess.run(
                [renderer, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(docx_path)],
                capture_output=True,
                timeout=90,
                check=False,
            )
            pdfs = list(Path(temp_dir).glob("*.pdf"))
            if result.returncode != 0 or len(pdfs) != 1:
                return {"status": "failed", "reason": "DOCX-to-PDF conversion failed"}
            pdf_bytes = pdfs[0].read_bytes()
            if not pdf_bytes.startswith(b"%PDF-") or len(pdf_bytes) < 512:
                return {"status": "failed", "reason": "rendered PDF is empty or malformed"}
            pages = len(re.findall(rb"/Type\s*/Page\b", pdf_bytes))
            if pages < 1:
                return {"status": "failed", "reason": "rendered PDF has no page objects"}
            return {"status": "passed", "pages": pages, "pdf_bytes": len(pdf_bytes)}
        except FileNotFoundError:
            return {"status": "not_run", "reason": "LibreOffice renderer is not installed"}
        except subprocess.TimeoutExpired:
            return {"status": "failed", "reason": "DOCX render timed out"}
        except Exception:
            return {"status": "failed", "reason": "DOCX render could not be inspected"}


def check_delivery(
    *,
    deliver_dir: Path,
    patent_title: str,
    final_markdown: Path,
    facts_ledger_path: Path,
    consistency_report: Path,
    ipr_report: Path,
    review_status_path: Path,
    manifest_path: Path,
    base_dir: Path,
    figure_delivery_mode: str,
    pre_export: bool = False,
) -> dict:
    checks: list[dict] = []
    not_run: list[str] = []
    deliver_root = deliver_dir.resolve()
    _check(checks, "delivery directory exists", "pass" if deliver_root.is_dir() else "fail")
    if not deliver_root.is_dir():
        return {
            "checks": checks, "checks_passed": False, "workflow_complete": False,
            "not_run": [], "errors": ["delivery directory was not found"],
        }

    declared_output = _manifest_value(manifest_path, "output_dir")
    bound = bool(declared_output) and os.path.normcase(str(Path(declared_output).resolve())) == os.path.normcase(str(deliver_root))
    _check(checks, "manifest output_dir matches delivery directory", "pass" if bound else "fail")

    if len(patent_title.strip()) > TITLE_MAX_CHARS:
        _check(checks, "title length", "fail", f"title exceeds {TITLE_MAX_CHARS} characters")
    else:
        _check(checks, "title length", "pass")
    manifest_title = _manifest_value(manifest_path, "final_title") or _manifest_value(manifest_path, "working_title")
    _check(checks, "delivery title matches manifest title", "pass" if manifest_title == patent_title else "fail")
    expected_docx = deliver_root / f"{patent_title}{DOCX_SUFFIX}"
    if not pre_export:
        _check(checks, "final DOCX exists at canonical path", "pass" if expected_docx.is_file() else "fail")

    final_md = final_markdown if final_markdown.is_absolute() else deliver_root / final_markdown
    md_inside = is_within(final_md, deliver_root)
    md_ok = md_inside and final_md.is_file() and final_md.stat().st_size > 0
    _check(checks, "final Markdown exists inside delivery directory", "pass" if md_ok else "fail")
    markdown = final_md.read_text(encoding="utf-8", errors="replace") if md_ok else ""
    title_heading = f"# {patent_title}"
    _check(checks, "final Markdown title matches manifest title", "pass" if title_heading in markdown.splitlines() else "fail")

    configured_mode = _manifest_value(manifest_path, "figure_delivery_mode") or "mermaid_only"
    mode_ok = figure_delivery_mode == configured_mode and figure_delivery_mode in {"mermaid_only", "mermaid_and_images"}
    _check(checks, "figure delivery mode matches manifest", "pass" if mode_ok else "fail")

    try:
        facts = json.loads(facts_ledger_path.read_text(encoding="utf-8"))
        facts_ok = isinstance(facts, dict) and facts.get("ledger_type") == "facts_ledger"
    except Exception:
        facts = {}
        facts_ok = False
    _check(checks, "facts ledger parses", "pass" if facts_ok else "fail")

    review_data: dict = {}
    try:
        review_data = json.loads(review_status_path.read_text(encoding="utf-8"))
        revision_required = (
            (base_dir / "artifacts/revision/phase_10_edit_plan.json").is_file()
            or (base_dir / "artifacts/revision/phase_10_structured_diff.json").is_file()
        )
        review_errors, review_state = validate_review_status(
            review_data,
            workspace=base_dir,
            require_revision=revision_required,
            output_dir=deliver_dir,
        )
        review_ok = not review_errors
    except Exception:
        review_errors = ["review status could not be parsed"]
        review_state = {"review_completed": False, "review_fresh": False}
        review_ok = False
    _check(checks, "explicit latest review status is complete and fresh", "pass" if review_ok else "fail")
    facts_reviewed = False
    if isinstance(review_data, dict) and isinstance(review_data.get("reviewed_materials"), dict):
        reviewed_files = review_data["reviewed_materials"].get("files")
        if isinstance(reviewed_files, dict):
            facts_reviewed = any(
                (Path(raw_path) if Path(raw_path).is_absolute() else base_dir / raw_path).resolve() == facts_ledger_path.resolve()
                and isinstance(expected, str)
                and facts_ledger_path.is_file()
                and sha256_file(facts_ledger_path).lower() == expected.lower()
                for raw_path, expected in reviewed_files.items()
                if isinstance(raw_path, str)
            )
    _check(checks, "reviewed material hash includes current facts ledger", "pass" if facts_reviewed else "fail")
    reviewed_md = str(
        (review_data.get("reviewed_materials") or {}).get("final_markdown_path", "")
    ) if isinstance(review_data, dict) else ""
    reviewed_md_path = Path(reviewed_md)
    if reviewed_md and not reviewed_md_path.is_absolute():
        reviewed_md_path = base_dir / reviewed_md_path
    _check(
        checks,
        "reviewed Markdown is the delivered Markdown",
        "pass" if md_ok and reviewed_md_path.resolve() == final_md.resolve() else "fail",
    )
    _check(checks, "consistency report exists", "pass" if consistency_report.is_file() else "fail")
    _check(checks, "IPR report exists", "pass" if ipr_report.is_file() else "fail")
    if review_errors:
        checks.append({"name": "review status validation", "result": "fail", "details": "; ".join(review_errors)})

    figures = facts.get("figure_registry", []) if isinstance(facts, dict) else []
    figure_mmd_sources: list[str] = []
    required_image_refs: set[str] = set()
    figure_errors = False
    for figure in figures if isinstance(figures, list) else []:
        if not isinstance(figure, dict):
            figure_errors = True
            continue
        artifacts = figure.get("artifacts") if isinstance(figure.get("artifacts"), dict) else {}
        mmd_ref = str(artifacts.get("mmd", "")).strip()
        mmd_path = _delivery_path(deliver_root, mmd_ref) if mmd_ref else None
        if not mmd_path or not mmd_path.is_file() or mmd_path.suffix.lower() != ".mmd":
            figure_errors = True
        else:
            mmd_text = mmd_path.read_text(encoding="utf-8", errors="replace")
            source_lines = [line.strip() for line in mmd_text.splitlines() if line.strip()]
            if not source_lines:
                figure_errors = True
            else:
                figure_mmd_sources.append(source_lines[0])
        if figure_delivery_mode == "mermaid_and_images":
            image_ref = str(artifacts.get("image", "")).strip()
            image_path = _delivery_path(deliver_root, image_ref) if image_ref else None
            if not image_path or not image_path.is_file():
                figure_errors = True
            else:
                required_image_refs.add(image_ref.replace("\\", "/"))
    _check(
        checks,
        "registered Mermaid sources exist in delivery directory",
        "pass" if bool(figures) and not figure_errors else "fail",
    )

    markdown_image_refs = [value.strip() for value in IMAGE_RE.findall(markdown)]
    image_paths_ok = True
    for image_ref in markdown_image_refs:
        image_path = _delivery_path(deliver_root, image_ref)
        if not image_path or not image_path.is_file():
            image_paths_ok = False
    if not required_image_refs.issubset({ref.replace("\\", "/") for ref in markdown_image_refs}):
        image_paths_ok = False
    if figure_delivery_mode == "mermaid_only" and markdown_image_refs:
        image_paths_ok = False
    _check(checks, "referenced image files exist and match configured mode", "pass" if image_paths_ok else "fail")

    docx_text = ""
    embedded_image_hashes: list[str] = []
    embedded_math_texts: list[str] = []
    if pre_export:
        docx_parse = "not_applicable"
    elif expected_docx.is_file():
        try:
            docx_text, embedded_image_hashes, embedded_math_texts = _docx_text_and_image_refs(expected_docx)
            docx_parse = "pass"
        except ImportError:
            docx_parse = "not_run"
            not_run.append("docx_text_inspection")
        except Exception:
            docx_parse = "fail"
    else:
        docx_parse = "fail"
    if not pre_export:
        _check(checks, "DOCX text and image relationships inspect", docx_parse)
        residues = "**" in docx_text or "$" in docx_text
        _check(checks, "DOCX has no raw Markdown or math delimiters", "fail" if residues else ("pass" if docx_parse == "pass" else "not_run"))
        docx_content_ok = (
            docx_parse == "pass"
            and patent_title in docx_text
            and all(line in docx_text for line in figure_mmd_sources)
        )
        _check(checks, "DOCX contains title and registered Mermaid sources", "pass" if docx_content_ok else ("fail" if docx_parse == "pass" else "not_run"))
        markdown_text_ok = docx_parse == "pass" and _markdown_text_present(markdown, docx_text)
        _check(
            checks,
            "DOCX retains visible Markdown text and table cells",
            "pass" if markdown_text_ok else ("fail" if docx_parse == "pass" else "not_run"),
            "DOCX is missing visible text from the reviewed Markdown" if docx_parse == "pass" and not markdown_text_ok else "",
        )
        if markdown_image_refs and docx_parse == "pass":
            expected_image_hashes = []
            for image_ref in markdown_image_refs:
                image_path = _delivery_path(deliver_root, image_ref)
                if image_path and image_path.is_file():
                    expected_image_hashes.append(sha256_file(image_path))
            embedded_covers_references = (
                len(expected_image_hashes) == len(markdown_image_refs)
                and not (Counter(expected_image_hashes) - Counter(embedded_image_hashes))
            )
            _check(checks, "DOCX embeds each Markdown-referenced image", "pass" if embedded_covers_references else "fail")
        elif docx_parse == "pass":
            _check(checks, "DOCX embeds each Markdown-referenced image", "pass")
        else:
            _check(checks, "DOCX embeds each Markdown-referenced image", "not_run")
            not_run.append("docx_image_relationship_check")

        try:
            expected_math_texts = [
                math_expression_text(expression)
                for expression in inline_math_expressions(markdown)
            ]
        except UnsupportedMathError:
            _check(
                checks, "DOCX inline math matches Markdown equations", "fail",
                "Markdown contains malformed or unsupported inline math syntax",
            )
        else:
            if docx_parse == "pass":
                math_matches = embedded_math_texts == expected_math_texts
                _check(
                    checks, "DOCX inline math matches Markdown equations",
                    "pass" if math_matches else "fail",
                    "" if math_matches else "DOCX OMML equations do not match Markdown inline math",
                )
            else:
                _check(checks, "DOCX inline math matches Markdown equations", "not_run")
                not_run.append("docx_math_check")

    evidence_ids = {
        str(evidence_id).strip()
        for feature in (facts.get("feature_registry", []) if isinstance(facts, dict) else [])
        if isinstance(feature, dict)
        for evidence_id in (feature.get("evidence_ids") or [])
        if isinstance(evidence_id, str) and evidence_id.strip()
    }
    citations_in_markdown = all(f"[[{evidence_id}]]" in markdown for evidence_id in evidence_ids)
    citations_in_docx = all(f"[[{evidence_id}]]" in docx_text for evidence_id in evidence_ids)
    citations_ok = citations_in_markdown and (pre_export or citations_in_docx)
    if evidence_ids:
        _check(checks, "feature evidence citations appear in Markdown and DOCX", "pass" if citations_ok else "fail")
    else:
        _check(checks, "feature evidence citations appear in Markdown and DOCX", "not_run", "facts ledger has no evidence IDs")
        not_run.append("citation_check")

    if pre_export:
        render = {"status": "not_applicable", "reason": "rendering is checked after the DOCX is exported"}
    else:
        render = _render_docx(expected_docx) if expected_docx.is_file() else {"status": "failed", "reason": "final DOCX missing"}
        _check(checks, "DOCX render health", render["status"], render.get("reason", ""))
        if render["status"] == "not_run":
            not_run.append("docx_render")

    failed = [check for check in checks if _normalized_check_status(check.get("result")) == "fail"]
    checks_passed = not failed
    workflow_complete = (
        not pre_export
        and checks_passed
        and not not_run
        and review_state.get("workflow_complete") is True
    )
    errors = [check["name"] for check in failed]
    export_allowed = bool(
        pre_export
        and checks_passed
        and not not_run
        and review_state.get("review_completed") is True
        and review_state.get("review_fresh") is True
        and review_state.get("workflow_complete") is True
    )
    return {
        "checks": checks,
        "checks_passed": checks_passed,
        "passed": checks_passed,
        "workflow_complete": workflow_complete,
        "export_allowed": export_allowed,
        "review_completed": review_state.get("review_completed") is True,
        "revision_validation": review_state.get("revision_validation", "not_run"),
        "render": render,
        "not_run": sorted(set(not_run)),
        "skipped": [],
        "errors": errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Delivery health gate")
    parser.add_argument("--deliver-dir", required=True)
    parser.add_argument("--patent-title", required=True)
    parser.add_argument("--final-markdown", required=True)
    parser.add_argument("--facts-ledger", required=True)
    parser.add_argument("--consistency-report", required=True)
    parser.add_argument("--ipr-report", required=True)
    parser.add_argument("--review-status", required=True)
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--figure-delivery-mode", choices=["mermaid_only", "mermaid_and_images"], required=True)
    parser.add_argument("--pre-export", action="store_true",
                        help="Check export prerequisites without requiring or validating the not-yet-generated DOCX")
    parser.add_argument("--out", default="artifacts/delivery/phase_11_delivery_health_report.json")
    parser.add_argument("--base-dir", default=".")
    args = parser.parse_args()
    report = check_delivery(
        deliver_dir=Path(args.deliver_dir),
        patent_title=args.patent_title,
        final_markdown=Path(args.final_markdown),
        facts_ledger_path=Path(args.facts_ledger),
        consistency_report=Path(args.consistency_report),
        ipr_report=Path(args.ipr_report),
        review_status_path=Path(args.review_status),
        manifest_path=Path(args.manifest),
        base_dir=Path(args.base_dir),
        figure_delivery_mode=args.figure_delivery_mode,
        pre_export=args.pre_export,
    )
    report.update({
        "doc_type": "delivery_health_report",
        "phase": "phase_11",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "deliver_dir": str(Path(args.deliver_dir).resolve()),
        "final_markdown_path": str(Path(args.final_markdown).resolve()),
        "final_docx_path": str((Path(args.deliver_dir) / f"{args.patent_title}{DOCX_SUFFIX}").resolve()),
        "figure_delivery_mode": args.figure_delivery_mode,
    })
    target = Path(args.out)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["checks_passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
