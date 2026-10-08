#!/usr/bin/env python3
"""Generate a disclosure DOCX without moving or deleting the source Markdown.

The requested output path is the actual output path. The DOCX is written to a
temporary sibling file and is published only after its package structure has
been checked. Existing outputs are preserved unless --overwrite is explicit.
Inline math supports letters, numbers, common Greek/operator commands,
subscripts, superscripts, grouped expressions, and fractions as native OMML.
Unsupported or malformed math is rejected instead of being silently altered.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from workflow_common import title_error

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, OSError, ValueError):
        pass

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt
except ImportError:
    Document = None


class UnsupportedMathError(ValueError):
    """An inline equation uses syntax outside the supported OMML subset."""


GREEK_COMMANDS = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
    "varepsilon": "ϵ", "zeta": "ζ", "eta": "η", "theta": "θ", "vartheta": "ϑ",
    "iota": "ι", "kappa": "κ", "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ",
    "omicron": "ο", "pi": "π", "varpi": "ϖ", "rho": "ρ", "varrho": "ϱ",
    "sigma": "σ", "varsigma": "ς", "tau": "τ", "upsilon": "υ", "phi": "φ",
    "varphi": "ϕ", "chi": "χ", "psi": "ψ", "omega": "ω", "Gamma": "Γ",
    "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ", "Pi": "Π",
    "Sigma": "Σ", "Upsilon": "Υ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω",
}
MATH_COMMANDS = {
    "times": "×", "cdot": "·", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥",
    "ne": "≠", "neq": "≠", "pm": "±", "to": "→", "rightarrow": "→",
    "in": "∈", "notin": "∉", "infty": "∞", "sum": "∑", "prod": "∏", "int": "∫",
}
MATH_OPERATOR_CHARS = set("+-=<>.,:;/!*|%×·≤≥≠±→∈∉∞∑∏∫")


def _font(run, name: str = "SimSun", size: int = 12, *, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.element.rPr.rFonts.set(qn("w:eastAsia"), name)


class _MathParser:
    def __init__(self, expression: str) -> None:
        self.expression = expression
        self.index = 0

    def parse(self) -> list[dict]:
        if not self.expression.strip():
            raise UnsupportedMathError("empty inline math is not supported")
        nodes = self._sequence()
        if self.index != len(self.expression):
            raise UnsupportedMathError("unsupported inline math syntax")
        return nodes

    def _skip_space(self) -> None:
        while self.index < len(self.expression) and self.expression[self.index].isspace():
            self.index += 1

    def _sequence(self, closing: str | None = None) -> list[dict]:
        nodes: list[dict] = []
        while self.index < len(self.expression):
            self._skip_space()
            if self.index >= len(self.expression):
                break
            char = self.expression[self.index]
            if closing and char == closing:
                self.index += 1
                return nodes
            if char in "})]":
                raise UnsupportedMathError("unbalanced inline math grouping")
            if char == "{":
                self.index += 1
                nodes.append({"kind": "group", "children": self._sequence("}")})
            elif char in "([":
                nodes.append(self._delimited())
            elif char == "\\":
                nodes.append(self._command())
            elif char in "^_":
                self._attach_script(nodes, "sup" if char == "^" else "sub")
            elif char.isalpha():
                nodes.append({"kind": "run", "text": char})
                self.index += 1
            elif char.isdigit():
                start = self.index
                while self.index < len(self.expression) and self.expression[self.index].isdigit():
                    self.index += 1
                if (
                    self.index + 1 < len(self.expression)
                    and self.expression[self.index] == "."
                    and self.expression[self.index + 1].isdigit()
                ):
                    self.index += 1
                    while self.index < len(self.expression) and self.expression[self.index].isdigit():
                        self.index += 1
                nodes.append({"kind": "run", "text": self.expression[start:self.index]})
            elif char in MATH_OPERATOR_CHARS:
                nodes.append({"kind": "run", "text": char})
                self.index += 1
            else:
                raise UnsupportedMathError("unsupported inline math syntax")
        if closing:
            raise UnsupportedMathError("unbalanced inline math grouping")
        return nodes

    def _delimited(self) -> dict:
        opening = self.expression[self.index]
        closing = ")" if opening == "(" else "]"
        self.index += 1
        return {
            "kind": "delimited", "opening": opening, "closing": closing,
            "children": self._sequence(closing),
        }

    def _command(self) -> dict:
        self.index += 1
        start = self.index
        while self.index < len(self.expression) and self.expression[self.index].isalpha():
            self.index += 1
        if start == self.index:
            raise UnsupportedMathError("unsupported inline math command")
        command = self.expression[start:self.index]
        if command == "frac":
            numerator = self._required_group()
            denominator = self._required_group()
            return {"kind": "fraction", "numerator": numerator, "denominator": denominator}
        if command in GREEK_COMMANDS:
            return {"kind": "run", "text": GREEK_COMMANDS[command]}
        if command in MATH_COMMANDS:
            return {"kind": "run", "text": MATH_COMMANDS[command]}
        raise UnsupportedMathError("unsupported inline math command")

    def _required_group(self) -> list[dict]:
        self._skip_space()
        if self.index >= len(self.expression) or self.expression[self.index] != "{":
            raise UnsupportedMathError("fraction arguments must be grouped with braces")
        self.index += 1
        nodes = self._sequence("}")
        if not nodes:
            raise UnsupportedMathError("empty grouped math arguments are not supported")
        return nodes

    def _script_argument(self) -> list[dict]:
        self._skip_space()
        if self.index >= len(self.expression):
            raise UnsupportedMathError("subscript or superscript is missing its argument")
        char = self.expression[self.index]
        if char == "{":
            return self._required_group()
        if char in "([":
            return [self._delimited()]
        if char == "\\":
            return [self._command()]
        if char.isalpha() or char in MATH_OPERATOR_CHARS:
            self.index += 1
            return [{"kind": "run", "text": char}]
        if char.isdigit():
            start = self.index
            while self.index < len(self.expression) and self.expression[self.index].isdigit():
                self.index += 1
            return [{"kind": "run", "text": self.expression[start:self.index]}]
        raise UnsupportedMathError("unsupported subscript or superscript argument")

    def _attach_script(self, nodes: list[dict], script_kind: str) -> None:
        self.index += 1
        if not nodes:
            raise UnsupportedMathError("subscript or superscript has no base")
        argument = self._script_argument()
        base = nodes.pop()
        if base.get("kind") == "script":
            if script_kind in base:
                raise UnsupportedMathError("duplicate subscript or superscript")
            base[script_kind] = argument
            nodes.append(base)
        else:
            nodes.append({"kind": "script", "base": base, script_kind: argument})


def _parse_math(expression: str) -> list[dict]:
    return _MathParser(expression).parse()


def _math_nodes_text(nodes: list[dict]) -> str:
    parts: list[str] = []
    for node in nodes:
        kind = node["kind"]
        if kind == "run":
            parts.append(node["text"])
        elif kind == "group":
            parts.append(_math_nodes_text(node["children"]))
        elif kind == "delimited":
            parts.extend((node["opening"], _math_nodes_text(node["children"]), node["closing"]))
        elif kind == "fraction":
            parts.append(_math_nodes_text(node["numerator"]) + _math_nodes_text(node["denominator"]))
        elif kind == "script":
            parts.append(_math_nodes_text([node["base"]]))
            parts.append(_math_nodes_text(node.get("sub", [])))
            parts.append(_math_nodes_text(node.get("sup", [])))
    return "".join(parts)


def inline_math_expressions(markdown: str) -> list[str]:
    """Return inline formulas outside fenced code blocks, rejecting malformed delimiters."""
    expressions: list[str] = []
    in_fence = False
    for line in markdown.splitlines():
        if line.strip().startswith(chr(96) * 3):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        if "$$" in line:
            raise UnsupportedMathError("display math delimiters are not supported")
        parts = line.split("$")
        if len(parts) % 2 == 0:
            raise UnsupportedMathError("inline math delimiters must be paired")
        expressions.extend(parts[1::2])
    return expressions


def math_expression_text(expression: str) -> str:
    return _math_nodes_text(_parse_math(expression))


def _math_run(text: str):
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def _append_math_nodes(parent, nodes: list[dict]) -> None:
    for node in nodes:
        kind = node["kind"]
        if kind == "run":
            parent.append(_math_run(node["text"]))
        elif kind == "group":
            _append_math_nodes(parent, node["children"])
        elif kind == "delimited":
            delimiter = OxmlElement("m:d")
            properties = OxmlElement("m:dPr")
            opening = OxmlElement("m:begChr")
            opening.set(qn("m:val"), node["opening"])
            closing = OxmlElement("m:endChr")
            closing.set(qn("m:val"), node["closing"])
            properties.extend((opening, closing))
            delimiter.append(properties)
            expression = OxmlElement("m:e")
            _append_math_nodes(expression, node["children"])
            delimiter.append(expression)
            parent.append(delimiter)
        elif kind == "fraction":
            fraction = OxmlElement("m:f")
            properties = OxmlElement("m:fPr")
            fraction_type = OxmlElement("m:type")
            fraction_type.set(qn("m:val"), "bar")
            properties.append(fraction_type)
            fraction.append(properties)
            numerator = OxmlElement("m:num")
            _append_math_nodes(numerator, node["numerator"])
            denominator = OxmlElement("m:den")
            _append_math_nodes(denominator, node["denominator"])
            fraction.extend((numerator, denominator))
            parent.append(fraction)
        elif kind == "script":
            if "sub" in node and "sup" in node:
                script = OxmlElement("m:sSubSup")
            elif "sub" in node:
                script = OxmlElement("m:sSub")
            else:
                script = OxmlElement("m:sSup")
            base = OxmlElement("m:e")
            _append_math_nodes(base, [node["base"]])
            script.append(base)
            if "sub" in node:
                sub = OxmlElement("m:sub")
                _append_math_nodes(sub, node["sub"])
                script.append(sub)
            if "sup" in node:
                sup = OxmlElement("m:sup")
                _append_math_nodes(sup, node["sup"])
                script.append(sup)
            parent.append(script)


def _add_math(p, expression: str) -> None:
    math = OxmlElement("m:oMath")
    _append_math_nodes(math, _parse_math(expression))
    p._p.append(math)


def _add_inline_segment(p, text: str, *, bold: bool) -> None:
    parts = text.split("$")
    if len(parts) % 2 == 0:
        raise UnsupportedMathError("inline math delimiters must be paired")
    for index, part in enumerate(parts):
        if not part:
            continue
        if index % 2:
            _add_math(p, part)
        else:
            run = p.add_run(part)
            _font(run, bold=bold)


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|") and not value.endswith("\\|"):
        value = value[:-1]
    cells: list[str] = []
    current: list[str] = []
    index = 0
    while index < len(value):
        char = value[index]
        if char == "\\" and index + 1 < len(value) and value[index + 1] == "|":
            current.append("|")
            index += 2
            continue
        if char == "|":
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
        index += 1
    cells.append("".join(current).strip())
    return cells


def _is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or "|" not in lines[index] or "|" not in lines[index + 1]:
        return False
    headers = _split_table_row(lines[index])
    separators = _split_table_row(lines[index + 1])
    return bool(headers) and len(headers) == len(separators) and all(
        re.fullmatch(r":?-{3,}:?", cell) for cell in separators
    )


def _add_table(doc, lines: list[str], index: int) -> int:
    headers = _split_table_row(lines[index])
    separators = _split_table_row(lines[index + 1])
    rows = [headers]
    next_index = index + 2
    while next_index < len(lines) and lines[next_index].strip() and "|" in lines[next_index]:
        cells = _split_table_row(lines[next_index])
        if len(cells) != len(headers):
            raise ValueError("Markdown table row has a different number of cells than its header")
        rows.append(cells)
        next_index += 1

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    alignments = []
    for marker in separators:
        if marker.startswith(":") and marker.endswith(":"):
            alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
        elif marker.endswith(":"):
            alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            alignments.append(WD_ALIGN_PARAGRAPH.LEFT)

    for row_index, cells in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for column, (cell, cell_text) in enumerate(zip(row.cells, cells)):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = alignments[column]
            _add_inline(paragraph, cell_text)
            if row_index == 0:
                for run in paragraph.runs:
                    _font(run, bold=True)
    return next_index


def _list_item(line: str) -> bool:
    value = line.strip()
    return bool(value) and any(re.match(pattern, value) for pattern in (
        r"^[a-zA-Z0-9]+[.)]",
        r"^\([a-zA-Z0-9]+\)",
        r"^[0-9]+[、.]",
    ))


def _add_inline(p, text: str) -> None:
    """Render paired Markdown bold and supported inline math without lossy fallback."""
    if "$$" in text:
        raise UnsupportedMathError("display math delimiters are not supported")
    if text.count("**") % 2:
        raise ValueError("Markdown bold delimiters must be paired")
    for bold_index, part in enumerate(text.split("**")):
        _add_inline_segment(p, part, bold=bool(bold_index % 2))


def _is_mermaid_open(line: str) -> bool:
    return line.strip().startswith(chr(96) * 3 + "mermaid")


def generate_docx(
    md_path: str | Path,
    docx_path: str | Path,
    *,
    overwrite: bool = False,
) -> dict | None:
    if Document is None:
        print("ERROR: python-docx is not installed", file=sys.stderr)
        return None
    md_file = Path(md_path)
    out = Path(docx_path)
    if not md_file.is_file():
        print("ERROR: input Markdown file does not exist", file=sys.stderr)
        return None
    if md_file.resolve() == out.resolve():
        print("ERROR: input Markdown and output DOCX paths must differ", file=sys.stderr)
        return None
    if out.is_symlink():
        raise ValueError("DOCX output path must not be a symbolic link")
    if out.exists() and not overwrite:
        raise FileExistsError("DOCX output already exists; pass --overwrite to replace it")

    content = md_file.read_text(encoding="utf-8")
    first_heading = next((line[2:].strip() for line in content.splitlines() if line.startswith("# ")), None)
    if first_heading and title_error(first_heading):
        print("ERROR: the first Markdown heading exceeds the 24-character title limit", file=sys.stderr)
        return None
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    images_embedded = 0
    images_missing = 0
    lines = content.splitlines()
    base_dir = md_file.parent
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# ") and not line.startswith("## "):
            p = doc.add_heading("", level=1)
            _add_inline(p, line[2:].strip())
            for run in p.runs:
                _font(run, "SimHei", 16, bold=True)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_heading("", level=2)
            _add_inline(p, line[3:].strip())
            for run in p.runs:
                _font(run, "SimHei", 14, bold=True)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_heading("", level=3)
            _add_inline(p, line[4:].strip())
            for run in p.runs:
                _font(run, "SimHei", 13, bold=True)
            i += 1
            continue

        if _is_table_start(lines, i):
            i = _add_table(doc, lines, i)
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if image_match:
            alt_text, image_ref = image_match.groups()
            image_path = (base_dir / image_ref).resolve()
            try:
                image_path.relative_to(base_dir.resolve())
                image_is_local = True
            except ValueError:
                image_is_local = False
            if image_is_local and image_path.is_file():
                try:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(image_path), width=Inches(5.5))
                    caption = doc.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_inline(caption, alt_text)
                    for run in caption.runs:
                        _font(run, size=10)
                    images_embedded += 1
                except Exception:
                    doc.add_paragraph(f"[Figure {alt_text} could not be embedded]")
                    images_missing += 1
            else:
                doc.add_paragraph(f"[Figure {alt_text} source image is missing]")
                images_missing += 1
            i += 1
            continue

        if _is_mermaid_open(line):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(chr(96) * 3):
                source_line = lines[i].rstrip()
                if source_line.strip():
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    run = p.add_run(source_line)
                    run.font.name = "Courier New"
                    run.font.size = Pt(7.5)
                i += 1
            i += 1
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            _add_inline(p, line[2:].strip())
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                _font(run, size=10)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue

        paragraph_lines = [line]
        j = i + 1
        if not _list_item(line):
            while (
                j < len(lines)
                and lines[j].strip()
                and not lines[j].startswith("#")
                and not lines[j].startswith(chr(96) * 3)
                and not lines[j].startswith(">")
                and not lines[j].startswith("![")
                and not _list_item(lines[j])
                and not _is_table_start(lines, j)
            ):
                paragraph_lines.append(lines[j])
                j += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0 if _list_item(line) else 0.74)
        p.paragraph_format.line_spacing = 1.5
        _add_inline(p, "".join(paragraph_lines))
        i = j

    fd, temp_name = tempfile.mkstemp(
        prefix=f".{out.stem}.", suffix=".tmp.docx", dir=out.parent
    )
    os.close(fd)
    temp_out = Path(temp_name)
    try:
        doc.save(str(temp_out))
        with zipfile.ZipFile(temp_out, "r") as package:
            if "word/document.xml" not in package.namelist() or package.testzip():
                raise ValueError("generated DOCX package failed structural verification")
        if overwrite:
            os.replace(temp_out, out)
        else:
            # Hard-link publication is atomic and fails if another process
            # created the destination after the initial existence check.
            os.link(temp_out, out)
    finally:
        try:
            temp_out.unlink(missing_ok=True)
        except OSError:
            pass

    return {
        "status": "generated",
        "output_path": str(out.resolve()),
        "images_embedded": images_embedded,
        "images_missing_count": images_missing,
        "source_markdown_preserved": md_file.is_file(),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate a local DOCX without replacing input Markdown")
    parser.add_argument("input", help="Input Markdown file")
    parser.add_argument("output", help="Output DOCX path")
    parser.add_argument("--overwrite", action="store_true", help="Replace an existing DOCX output atomically")
    args = parser.parse_args(sys.argv[1:] if argv is None else argv)
    try:
        result = generate_docx(args.input, args.output, overwrite=args.overwrite)
    except Exception:
        print(json.dumps({"status": "failed", "error": "DOCX export failed"}, ensure_ascii=False))
        return 1
    if result is None:
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
